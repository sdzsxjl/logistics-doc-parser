"""
Excel导出模块 - 将解析结果导出为格式化的Excel文件
"""
import io
import os
from datetime import datetime
from typing import List
from openpyxl import Workbook
from openpyxl.styles import (
    Font, PatternFill, Alignment, Border, Side, NamedStyle
)
from openpyxl.utils import get_column_letter
import pandas as pd

from .models import ShippingOrder


class ExcelExporter:
    """Excel导出器"""

    # 颜色定义
    HEADER_FILL = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    HEADER_FONT = Font(name="微软雅黑", size=11, bold=True, color="FFFFFF")
    DATA_FONT = Font(name="微软雅黑", size=10)
    TITLE_FONT = Font(name="微软雅黑", size=14, bold=True, color="1F4E79")
    ALTERNATE_FILL = PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid")
    WARNING_FILL = PatternFill(start_color="FFF3CD", end_color="FFF3CD", fill_type="solid")
    WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

    THIN_BORDER = Border(
        left=Side(style="thin", color="D0D0D0"),
        right=Side(style="thin", color="D0D0D0"),
        top=Side(style="thin", color="D0D0D0"),
        bottom=Side(style="thin", color="D0D0D0"),
    )

    # Excel列定义
    COLUMNS = [
        ("运单号", 18),
        ("订单号", 18),
        ("发货日期", 12),
        ("重量", 10),
        ("件数", 8),
        ("运输方式", 12),
        ("承运商", 12),
        ("寄件人姓名", 15),
        ("寄件人电话", 15),
        ("寄件人地址", 40),
        ("寄件人省份", 10),
        ("寄件人城市", 12),
        ("寄件人区县", 12),
        ("收件人姓名", 15),
        ("收件人电话", 15),
        ("收件人地址", 40),
        ("收件人省份", 10),
        ("收件人城市", 12),
        ("收件人区县", 12),
        ("申报价值", 12),
        ("付款方式", 10),
        ("预计送达", 12),
        ("备注", 20),
        ("置信度", 8),
        ("来源文件", 25),
    ]

    def export(self, orders: List[ShippingOrder], filename_prefix: str = "物流信息") -> bytes:
        """
        将运单列表导出为Excel文件的字节数据

        Args:
            orders: 运单列表
            filename_prefix: 文件名前缀

        Returns:
            Excel文件的字节数据（可直接供下载）
        """
        wb = Workbook()

        # Sheet 1: 运单数据
        self._create_shipments_sheet(wb, orders)

        # Sheet 2: 统计汇总
        if orders:
            self._create_summary_sheet(wb, orders)

        # 保存到BytesIO
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    def export_to_file(self, orders: List[ShippingOrder], output_path: str):
        """导出到文件"""
        data = self.export(orders)
        os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
        with open(output_path, "wb") as f:
            f.write(data)

    def _create_shipments_sheet(self, wb: Workbook, orders: List[ShippingOrder]):
        """创建运单数据Sheet"""
        ws = wb.active
        ws.title = "运单数据"

        # 标题行
        ws.merge_cells("A1:Y1")
        title_cell = ws["A1"]
        title_cell.value = "物流运单信息提取结果"
        title_cell.font = self.TITLE_FONT
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 35

        # 副标题
        ws.merge_cells("A2:Y2")
        ws["A2"].value = f"导出时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}  |  共 {len(orders)} 条记录"
        ws["A2"].font = Font(name="微软雅黑", size=9, color="666666")
        ws["A2"].alignment = Alignment(horizontal="center")
        ws.row_dimensions[2].height = 22

        # 表头（第4行）
        header_row = 4
        for col_idx, (col_name, col_width) in enumerate(self.COLUMNS, 1):
            cell = ws.cell(row=header_row, column=col_idx)
            cell.value = col_name
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = self.THIN_BORDER
            ws.column_dimensions[get_column_letter(col_idx)].width = col_width

        ws.row_dimensions[header_row].height = 30

        # 冻结表头
        ws.freeze_panes = f"A{header_row + 1}"

        # 数据行
        for row_idx, order in enumerate(orders):
            excel_row = header_row + 1 + row_idx
            flat = order.to_flat_dict()

            for col_idx, (col_name, _) in enumerate(self.COLUMNS, 1):
                cell = ws.cell(row=excel_row, column=col_idx)
                value = flat.get(col_name, "")

                # 置信度格式化
                if col_name == "置信度" and isinstance(value, (int, float)):
                    cell.value = value
                    cell.number_format = "0.0%"
                else:
                    cell.value = value if value else ""

                cell.font = self.DATA_FONT
                cell.border = self.THIN_BORDER
                cell.alignment = Alignment(vertical="center", wrap_text=True)

                # 交替行颜色
                if row_idx % 2 == 1:
                    cell.fill = self.ALTERNATE_FILL

                # 低置信度标记
                if col_name == "置信度" and isinstance(value, (int, float)) and value < 0.6:
                    cell.fill = self.WARNING_FILL

            # 行高
            ws.row_dimensions[excel_row].height = 22

        # 自动筛选
        last_col = get_column_letter(len(self.COLUMNS))
        ws.auto_filter.ref = f"A{header_row}:{last_col}{header_row + len(orders)}"

    def _create_summary_sheet(self, wb: Workbook, orders: List[ShippingOrder]):
        """创建统计汇总Sheet"""
        ws = wb.create_sheet("统计汇总")

        # 标题
        ws.merge_cells("A1:C1")
        ws["A1"].value = "解析统计汇总"
        ws["A1"].font = self.TITLE_FONT
        ws.row_dimensions[1].height = 30

        # 统计数据
        stats = [
            ("总记录数", len(orders)),
            ("AI提取数", sum(1 for o in orders if o.extraction_method == "ai")),
            ("正则提取数", sum(1 for o in orders if o.extraction_method == "regex")),
            ("有运单号的记录", sum(1 for o in orders if o.tracking_number)),
            ("有寄件人的记录", sum(1 for o in orders if o.sender.name)),
            ("有收件人的记录", sum(1 for o in orders if o.recipient.name)),
            ("平均置信度", f"{sum(o.extraction_confidence for o in orders) / len(orders):.1%}" if orders else "N/A"),
        ]

        for i, (label, value) in enumerate(stats, 3):
            ws.cell(row=i, column=1, value=label).font = Font(name="微软雅黑", size=11, bold=True)
            ws.cell(row=i, column=2, value=value).font = Font(name="微软雅黑", size=11)
            ws.row_dimensions[i].height = 24

        ws.column_dimensions["A"].width = 18
        ws.column_dimensions["B"].width = 15
        ws.column_dimensions["C"].width = 15

        # 承运商分布
        carrier_count = {}
        for o in orders:
            if o.carrier:
                carrier_count[o.carrier] = carrier_count.get(o.carrier, 0) + 1

        if carrier_count:
            row = len(stats) + 5
            ws.cell(row=row, column=1, value="承运商分布").font = Font(
                name="微软雅黑", size=11, bold=True, color="1F4E79"
            )
            row += 1
            for carrier, count in sorted(carrier_count.items(), key=lambda x: x[1], reverse=True):
                ws.cell(row=row, column=1, value=carrier).font = Font(name="微软雅黑", size=10)
                ws.cell(row=row, column=2, value=count).font = Font(name="微软雅黑", size=10)
                row += 1


def export_to_excel(orders: List[ShippingOrder], output_path: str = ""):
    """便捷函数：导出运单到Excel"""
    exporter = ExcelExporter()

    if output_path:
        exporter.export_to_file(orders, output_path)
        return output_path
    else:
        return exporter.export(orders)


def orders_to_dataframe(orders: List[ShippingOrder]) -> "pd.DataFrame":
    """将运单列表转为pandas DataFrame"""
    if not orders:
        return pd.DataFrame()
    rows = [order.to_flat_dict() for order in orders]
    return pd.DataFrame(rows)
